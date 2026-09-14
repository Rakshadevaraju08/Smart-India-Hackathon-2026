import os
import pptx
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

def assemble_final_deck():
    base_dir = r"c:\Users\Gagan K S\Documents\SIH"
    src_pptx = os.path.join(base_dir, "SIH2026_Idea_Presentation_Updated.pptx")
    out_pptx = os.path.join(base_dir, "SIH2026_Idea_Presentation_FINAL.pptx")
    out_pdf = os.path.join(base_dir, "SIH2026_Idea_Presentation_FINAL.pdf")

    prs = pptx.Presentation(src_pptx)
    print(f"Loaded presentation with {len(prs.slides)} slides.")

    # -------------------------------------------------------------
    # SLIDE 1: Title & Details
    # -------------------------------------------------------------
    s1 = prs.slides[0]
    for shape in s1.shapes:
        if shape.name == "Subtitle 3" and shape.has_text_frame:
            shape.text_frame.clear()
            p = shape.text_frame.paragraphs[0]
            p.text = "Team Kairos"
            p.font.name = "Arial"
            p.font.size = Pt(40)
            p.font.bold = True
            p.font.color.rgb = RGBColor(15, 23, 42)
            p.alignment = PP_ALIGN.CENTER
        
        elif shape.name == "TextBox 9" and shape.has_text_frame:
            shape.text_frame.clear()
            details = [
                ("Problem Statement ID", "26085"),
                ("Problem Statement Title", "Urban Flood Nowcasting System (Coupled Drainage & Rainfall)"),
                ("Organization", "Ministry of Earth Sciences (MoES) / NCMRWF"),
                ("Theme", "Disaster Management"),
                ("PS Category", "Software"),
                ("Team ID", "SIH2026/26085/KAIROS"),
                ("Team Name", "Team Kairos"),
            ]
            for idx, (label, val) in enumerate(details):
                p = shape.text_frame.add_paragraph() if idx > 0 else shape.text_frame.paragraphs[0]
                p.space_after = Pt(8)
                r_lbl = p.add_run()
                r_lbl.text = f"{label}: "
                r_lbl.font.name = "Calibri"
                r_lbl.font.size = Pt(17)
                r_lbl.font.bold = True
                r_lbl.font.color.rgb = RGBColor(30, 41, 59)

                r_val = p.add_run()
                r_val.text = val
                r_val.font.name = "Calibri"
                r_val.font.size = Pt(17)
                r_val.font.bold = False
                r_val.font.color.rgb = RGBColor(15, 23, 42)

    # Helper function to remove a shape by name
    def remove_shape_by_name(slide, name):
        for shp in list(slide.shapes):
            if shp.name == name:
                sp = shp._element
                sp.getparent().remove(sp)

    # Helper to set Oval team name
    def update_oval_team_name(slide):
        for shp in slide.shapes:
            if "Oval" in shp.name and shp.has_text_frame:
                shp.text_frame.clear()
                p = shp.text_frame.paragraphs[0]
                p.text = "Team Kairos"
                p.font.name = "Arial"
                p.font.size = Pt(11)
                p.font.bold = True
                p.font.color.rgb = RGBColor(255, 255, 255)
                p.alignment = PP_ALIGN.CENTER

    # -------------------------------------------------------------
    # SLIDE 2: Proposed Solution
    # -------------------------------------------------------------
    s2 = prs.slides[1]
    update_oval_team_name(s2)
    for shp in s2.shapes:
        if shp.name == "Title 1" and shp.has_text_frame:
            p = shp.text_frame.paragraphs[0]
            p.text = "PROPOSED SOLUTION & NOVELTY"
            p.font.name = "Calibri"
            p.font.size = Pt(28)
            p.font.bold = True
            p.font.color.rgb = RGBColor(15, 23, 42)
    remove_shape_by_name(s2, "TextBox 8")
    s2_img = os.path.join(base_dir, "slide2_problem_solution_tiranga_glare.png")
    s2.shapes.add_picture(s2_img, Inches(0.40), Inches(1.35), Inches(12.53), Inches(5.55))
    print("Slide 2 configured with infographic.")

    # -------------------------------------------------------------
    # SLIDE 3: Technical Approach
    # -------------------------------------------------------------
    s3 = prs.slides[2]
    update_oval_team_name(s3)
    for shp in s3.shapes:
        if shp.name == "Title 1" and shp.has_text_frame:
            p = shp.text_frame.paragraphs[0]
            p.text = "TECHNICAL APPROACH & ARCHITECTURE"
            p.font.name = "Calibri"
            p.font.size = Pt(28)
            p.font.bold = True
            p.font.color.rgb = RGBColor(15, 23, 42)
    print("Slide 3 title & team name verified.")

    # -------------------------------------------------------------
    # SLIDE 4: Feasibility & Viability
    # -------------------------------------------------------------
    s4 = prs.slides[3]
    update_oval_team_name(s4)
    for shp in s4.shapes:
        if shp.name == "Title 1" and shp.has_text_frame:
            p = shp.text_frame.paragraphs[0]
            p.text = "FEASIBILITY AND VIABILITY"
            p.font.name = "Calibri"
            p.font.size = Pt(28)
            p.font.bold = True
            p.font.color.rgb = RGBColor(15, 23, 42)
    print("Slide 4 title & team name verified.")

    # -------------------------------------------------------------
    # SLIDE 5: Impact & Benefits
    # -------------------------------------------------------------
    s5 = prs.slides[4]
    update_oval_team_name(s5)
    for shp in s5.shapes:
        if shp.name == "Title 1" and shp.has_text_frame:
            p = shp.text_frame.paragraphs[0]
            p.text = "IMPACT, SOCIAL BENEFITS & COMMERCIAL POTENTIAL"
            p.font.name = "Calibri"
            p.font.size = Pt(26)
            p.font.bold = True
            p.font.color.rgb = RGBColor(15, 23, 42)
    remove_shape_by_name(s5, "TextBox 8")
    s5_img = os.path.join(base_dir, "slide5_impact_benefits_tiranga_glare.png")
    s5.shapes.add_picture(s5_img, Inches(0.40), Inches(1.35), Inches(12.53), Inches(5.55))
    print("Slide 5 configured with infographic.")

    # -------------------------------------------------------------
    # SLIDE 6: Research & References
    # -------------------------------------------------------------
    s6 = prs.slides[5]
    update_oval_team_name(s6)
    for shp in s6.shapes:
        if shp.name == "Title 1" and shp.has_text_frame:
            p = shp.text_frame.paragraphs[0]
            p.text = "RESEARCH, REFERENCES & VALIDATION PROOF"
            p.font.name = "Calibri"
            p.font.size = Pt(26)
            p.font.bold = True
            p.font.color.rgb = RGBColor(15, 23, 42)
    remove_shape_by_name(s6, "TextBox 8")
    s6_img = os.path.join(base_dir, "slide6_research_references_tiranga_glare.png")
    s6.shapes.add_picture(s6_img, Inches(0.40), Inches(1.35), Inches(12.53), Inches(5.55))
    print("Slide 6 configured with infographic.")

    # -------------------------------------------------------------
    # SLIDE 7: DELETE SLIDE 7
    # -------------------------------------------------------------
    if len(prs.slides) >= 7:
        rId = prs.slides._sldIdLst[6].rId
        prs.part.drop_rel(rId)
        del prs.slides._sldIdLst[6]
        print("Slide 7 (Important Instructions) successfully deleted.")

    print(f"Final presentation has {len(prs.slides)} slides.")
    prs.save(out_pptx)
    print(f"Saved: {out_pptx}")

    # -------------------------------------------------------------
    # EXPORT TO PDF VIA POWERPOINT COM
    # -------------------------------------------------------------
    print("Exporting cleanly to PDF via PowerPoint COM...")
    try:
        import win32com.client
        import time
        ppt_app = win32com.client.DispatchEx("PowerPoint.Application")
        presentation = ppt_app.Presentations.Open(os.path.abspath(out_pptx), WithWindow=False)
        presentation.SaveAs(os.path.abspath(out_pdf), 32)
        presentation.Close()
        ppt_app.Quit()
        print(f"Successfully exported PDF: {out_pdf} ({os.path.getsize(out_pdf)} bytes)")
    except Exception as e:
        print(f"PowerPoint COM export error: {e}")

if __name__ == "__main__":
    assemble_final_deck()
